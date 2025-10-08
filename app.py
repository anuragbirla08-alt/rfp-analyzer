# app.py - Enhanced RFP Analyzer with Two-Step Process and Excel Append

import json
import os
import re
import mimetypes
import zipfile
import time
from io import BytesIO
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Tuple, Union
from functools import wraps

import pandas as pd
import requests
import streamlit as st
from docx import Document
import fitz  # PyMuPDF
from google import genai
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment, PatternFill
from datetime import datetime

# --- Page Configuration ---
st.set_page_config(
    page_title="RFP Analysis Tool",
    page_icon="📄",
    layout="wide"
)

st.title("📄 RFP Analysis Automation Tool")
st.write(
    "Upload government tender PDFs to extract key information in two steps: "
    "first analyze the main RFP, then optionally analyze supporting documents."
)

# --- Constants and Defaults ---
API_KEY_PATH = Path.home() / ".rfp_analyzer_api_key"
API_KEYS_PATH = Path.home() / ".rfp_analyzer_api_keys.json"  # Multiple keys storage
DEFAULT_PROMPT = """You are an expert government tender analyst. Read the full RFP carefully and return a single JSON object only.
Capture the following baseline fields as strings (join lists with \n where needed).
- bid_number
- start_date
- end_date
- city
- organisation_name
- department_name
- period_of_contract
- estimated_bid_amount
- category_or_type
- scope_of_work (array of bullet points that retain the concrete activities, deliverables, frequencies, timelines, and compliance responsibilities. Keep one bullet per obligation and preserve numbering or lettering.)
- eligibility_criteria (array of bullet points that spell out turnover values, certifications, manpower mix, past experience, geography, registrations, or any disqualifications verbatim. Include thresholds even if repeated elsewhere.)
- delivery_or_execution_mode (capture onsite/offsite/hybrid requirements and any visit frequency expectations)
- contact_details
- minimum_turnover_requirement
- experience_requirement
- earnest_money_deposit
- bid_security
- important_dates (array; include bid opening, pre-bid, clarification deadlines, contract award, etc.)

Additionally:
- Place every other commercial, technical, submission, penalty, payment, staffing, document checklist, or evaluation requirement inside an "additional_fields" object using descriptive keys.
- Do not omit financial numbers such as turnover thresholds, bid values, security deposits, or penalties even if they appear in multiple sections.
- Preserve numbering and sub-bullets exactly as written within list items.
- Respond with valid JSON only. Do not include markdown, commentary, or triple backticks.
- If a field is absent, set it to "Not Found" (or an empty array for list fields)."""

DOCUMENT_FILE_EXTENSIONS = {".pdf", ".doc", ".docx", ".xls", ".xlsx", ".zip"}
DOCUMENT_LINK_KEYWORDS: Tuple[str, ...] = (
    "sow", "scope", "atc", "additional", "annex", "corrigendum", "tender", "document", "attachment",
)

# Initialize session state
if "primary_analysis_results" not in st.session_state:
    st.session_state["primary_analysis_results"] = []
if "supporting_analysis_results" not in st.session_state:
    st.session_state["supporting_analysis_results"] = []
if "extracted_links" not in st.session_state:
    st.session_state["extracted_links"] = {}
if "attachment_cache" not in st.session_state:
    st.session_state["attachment_cache"] = {}
if "current_step" not in st.session_state:
    st.session_state["current_step"] = 1
if "existing_excel_data" not in st.session_state:
    st.session_state["existing_excel_data"] = None
if "existing_excel_filename" not in st.session_state:
    st.session_state["existing_excel_filename"] = None
if "api_keys_list" not in st.session_state:
    st.session_state["api_keys_list"] = []
if "current_api_key_index" not in st.session_state:
    st.session_state["current_api_key_index"] = 0
if "failed_api_keys" not in st.session_state:
    st.session_state["failed_api_keys"] = set()

# --- Utility Functions ---
def load_api_keys() -> List[str]:
    """Load multiple API keys from storage."""
    keys = []
    
    # Try loading from multi-key storage
    if API_KEYS_PATH.exists():
        try:
            with open(API_KEYS_PATH, 'r') as f:
                data = json.load(f)
                keys = data.get('keys', [])
        except:
            pass
    
    # Fallback to single key storage
    if not keys and API_KEY_PATH.exists():
        try:
            single_key = API_KEY_PATH.read_text(encoding="utf-8").strip()
            if single_key:
                keys = [single_key]
        except:
            pass
    
    # Check environment variable
    env_key = os.getenv("GENAI_API_KEY")
    if env_key and env_key.strip() not in keys:
        keys.insert(0, env_key.strip())
    
    return keys

def save_api_keys(keys: List[str]) -> None:
    """Save multiple API keys to storage."""
    keys = [k.strip() for k in keys if k.strip()]
    if not keys:
        st.error("At least one API key is required.")
        return
    
    try:
        with open(API_KEYS_PATH, 'w') as f:
            json.dump({'keys': keys}, f)
        os.chmod(API_KEYS_PATH, 0o600)
        st.session_state["api_keys_list"] = keys
        st.session_state["current_api_key_index"] = 0
        st.session_state["failed_api_keys"] = set()
        st.success(f"✅ Saved {len(keys)} API key(s). They will rotate automatically when limits are hit.")
    except Exception as e:
        st.error(f"Failed to save API keys: {e}")

def get_next_api_key() -> Optional[str]:
    """Get the next available API key, rotating through the list."""
    keys = st.session_state.get("api_keys_list", [])
    if not keys:
        return None
    
    failed_keys = st.session_state.get("failed_api_keys", set())
    available_keys = [k for k in keys if k not in failed_keys]
    
    if not available_keys:
        # All keys failed, reset and try again
        st.warning("⚠️ All API keys exhausted. Resetting rotation...")
        st.session_state["failed_api_keys"] = set()
        available_keys = keys
    
    current_index = st.session_state.get("current_api_key_index", 0)
    
    # Find next available key
    for i in range(len(keys)):
        idx = (current_index + i) % len(keys)
        if keys[idx] in available_keys:
            st.session_state["current_api_key_index"] = idx
            return keys[idx]
    
    return keys[0] if keys else None

def mark_api_key_failed(api_key: str) -> None:
    """Mark an API key as failed (rate limited)."""
    if "failed_api_keys" not in st.session_state:
        st.session_state["failed_api_keys"] = set()
    st.session_state["failed_api_keys"].add(api_key)
    
    # Move to next key
    keys = st.session_state.get("api_keys_list", [])
    if keys and api_key in keys:
        current_idx = keys.index(api_key)
        st.session_state["current_api_key_index"] = (current_idx + 1) % len(keys)

def load_api_key() -> Optional[str]:
    """Load the first available API key (for backward compatibility)."""
    keys = load_api_keys()
    return keys[0] if keys else None

def get_client(api_key: str) -> genai.Client:
    return genai.Client(api_key=api_key)

def extract_text_from_pdf(pdf_bytes: bytes) -> Optional[str]:
    """Extract text from PDF bytes."""
    try:
        pdf_document = fitz.open(stream=pdf_bytes, filetype="pdf")
        page_count = len(pdf_document)
        st.info(f"📖 PDF has {page_count} pages")
        
        text_parts = []
        empty_pages = 0
        
        for page_num, page in enumerate(pdf_document, 1):
            page_text = page.get_text()
            if page_text.strip():
                text_parts.append(page_text)
            else:
                empty_pages += 1
                st.warning(f"⚠️ Page {page_num} appears empty (might be scanned image)")
        
        pdf_document.close()
        
        if empty_pages == page_count:
            st.error(f"❌ All {page_count} pages are empty - this is likely a scanned/image-based PDF")
            st.info("💡 Tip: Use OCR software to convert scanned PDFs to text-based PDFs")
            return None
        
        if empty_pages > 0:
            st.warning(f"⚠️ {empty_pages}/{page_count} pages were empty")
        
        full_text = "\n".join(text_parts)
        
        if not full_text.strip():
            st.error("❌ No text could be extracted - PDF might contain only images")
            return None
            
        return full_text
        
    except Exception as exc:
        st.error(f"❌ Error reading PDF: {exc}")
        import traceback
        st.error(f"Traceback: {traceback.format_exc()}")
        return None

def extract_text_from_docx(docx_bytes: bytes) -> Optional[str]:
    try:
        document = Document(BytesIO(docx_bytes))
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        return text if text.strip() else None
    except Exception as exc:
        st.error(f"Unable to read DOCX file: {exc}")
        return None

def extract_attachments_from_pdf(pdf_bytes: bytes) -> List[Tuple[str, bytes, Optional[str]]]:
    """Return embedded files inside the PDF (name, data, description)."""
    attachments: List[Tuple[str, bytes, Optional[str]]] = []
    try:
        pdf_document = fitz.open(stream=pdf_bytes, filetype="pdf")
        getter = getattr(pdf_document, "embeddedFileGet", None)
        if callable(getter):
            count = getattr(pdf_document, "embeddedFileCount", 0) or 0
            names = list(getattr(pdf_document, "embeddedFileNames", lambda: [])())
            for index in range(count):
                info = None
                try:
                    info = getter(index)
                except TypeError:
                    if index < len(names):
                        info = getter(names[index])
                if not isinstance(info, dict):
                    continue
                name = info.get("name") or (names[index] if index < len(names) else f"attachment_{index + 1}")
                file_bytes = info.get("file", b"")
                description = info.get("desc")
                if file_bytes:
                    attachments.append((name, file_bytes, description))
        elif hasattr(pdf_document, "embeddedFileNames"):
            for name in pdf_document.embeddedFileNames():
                info = pdf_document.embeddedFileGet(name)
                file_bytes = info.get("file", b"") if isinstance(info, dict) else b""
                if not file_bytes:
                    continue
                attachments.append((info.get("name", name), file_bytes, info.get("desc") if isinstance(info, dict) else None))
        pdf_document.close()
    except Exception as exc:
        st.warning(f"Unable to extract embedded files: {exc}")
    return attachments

def extract_embedded_links_from_pdf(pdf_bytes: bytes) -> List[Tuple[str, str, str, str]]:
    """Extract embedded hyperlinks from PDF annotations (url, link_text, page_number, label)."""
    embedded_links = []
    seen = set()
    
    try:
        pdf_document = fitz.open(stream=pdf_bytes, filetype="pdf")
        for page_num, page in enumerate(pdf_document, start=1):
            links = page.get_links()
            
            for link in links:
                if link.get("kind") == 2:  # URI link
                    uri = link.get("uri", "")
                    if uri and uri not in seen:
                        seen.add(uri)
                        rect = link.get("from")
                        link_text = "Click here"
                        label = "Document"
                        
                        if rect:
                            try:
                                text_instances = page.get_text("words")
                                nearby_words = []
                                
                                for word in text_instances:
                                    word_rect = fitz.Rect(word[:4])
                                    if word_rect.intersects(rect) or word_rect.x0 <= rect.x1 + 100:
                                        nearby_words.append(word[4])
                                
                                if nearby_words:
                                    context = " ".join(nearby_words[:15])
                                    context_lower = context.lower()
                                    
                                    if "atc" in context_lower:
                                        label = "ATC Document"
                                    elif "sow" in context_lower or "scope of work" in context_lower:
                                        label = "Scope of Work (SOW)"
                                    elif "corrigendum" in context_lower:
                                        label = "Corrigendum"
                                    elif "annex" in context_lower or "annexure" in context_lower:
                                        label = "Annexure"
                                    elif "tender" in context_lower:
                                        label = "Tender Document"
                                    
                                    link_text = context[:60]
                            except:
                                pass
                        
                        embedded_links.append((uri, link_text, f"Page {page_num}", label))
        pdf_document.close()
    except Exception as exc:
        st.warning(f"Unable to extract embedded links: {exc}")
    
    return embedded_links

def find_reference_links(text: str) -> List[Tuple[str, str]]:
    """Extract plain text URLs and categorize them."""
    url_pattern = re.compile(r"https?://[^\s>]+", re.IGNORECASE)
    links = []
    seen = set()
    
    for match in url_pattern.findall(text):
        cleaned = match.rstrip(').,;"\'\n')
        if cleaned in seen:
            continue
        seen.add(cleaned)
        
        lowered = cleaned.lower()
        if any(keyword in lowered for keyword in DOCUMENT_LINK_KEYWORDS):
            category = "Related Document"
        elif Path(lowered).suffix in DOCUMENT_FILE_EXTENSIONS:
            category = "Other Document"
        else:
            category = "General Link"
        
        links.append((cleaned, category))
    
    return links

def categorize_link(url: str, context: str = "", label: str = "") -> str:
    """Categorize a URL based on its content, context, and label."""
    lowered = url.lower()
    context_lower = context.lower()
    label_lower = label.lower()
    
    if "atc" in label_lower:
        return "ATC Document"
    if "sow" in label_lower or "scope of work" in label_lower:
        return "SOW Document"
    if "corrigendum" in label_lower:
        return "Corrigendum"
    if "annex" in label_lower:
        return "Annexure"
    if "atc" in lowered or "atc" in context_lower:
        return "ATC Document"
    if "sow" in lowered or "scope" in lowered or "sow" in context_lower or "scope of work" in context_lower:
        return "SOW Document"
    if "corrigendum" in lowered or "corrigendum" in context_lower:
        return "Corrigendum"
    if any(keyword in lowered for keyword in DOCUMENT_LINK_KEYWORDS):
        return "Related Document"
    elif Path(lowered).suffix in DOCUMENT_FILE_EXTENSIONS:
        return "Other Document"
    else:
        return "General Link"

def fetch_external_document(url: str) -> Optional[Tuple[str, bytes, Optional[str]]]:
    """Download document from URL."""
    try:
        response = requests.get(url, timeout=30)
        response.raise_for_status()
        filename = url.split("/")[-1] or "downloaded_file"
        if "content-disposition" in response.headers:
            disposition = response.headers["content-disposition"]
            match = re.search(r'filename="?([^";]+)"?', disposition)
            if match:
                filename = match.group(1)
        content_type = response.headers.get("content-type")
        return filename, response.content, content_type
    except requests.RequestException as exc:
        st.error(f"Failed to download {url}: {exc}")
        return None

# Add retry decorator for API calls
def retry_on_failure(max_attempts=3, delay=2):
    """Decorator to retry function on failure with exponential backoff."""
    def decorator(func):
        @wraps(func)
        def wrapper(*args, **kwargs):
            for attempt in range(max_attempts):
                try:
                    result = func(*args, **kwargs)
                    if result:  # If we got a valid result, return it
                        return result
                    if attempt < max_attempts - 1:
                        wait_time = delay * (2 ** attempt)
                        st.warning(f"Attempt {attempt + 1} failed. Retrying in {wait_time}s...")
                        time.sleep(wait_time)
                except Exception as e:
                    if attempt < max_attempts - 1:
                        wait_time = delay * (2 ** attempt)
                        st.warning(f"Error on attempt {attempt + 1}: {str(e)}. Retrying in {wait_time}s...")
                        time.sleep(wait_time)
                    else:
                        st.error(f"Failed after {max_attempts} attempts: {str(e)}")
                        raise
            return None
        return wrapper
    return decorator

def run_model(prompt_template: str, text_content: str, api_key: str, max_retries: int = 3, allow_key_rotation: bool = True) -> str:
    """Run the AI model with proper error handling, rate limiting, and key rotation."""
    prompt = f"{prompt_template.strip()}\n\n---\n{text_content[:50000]}\n---"
    
    for attempt in range(max_retries):
        try:
            client = get_client(api_key)
            response = client.models.generate_content(
                model="gemini-2.0-flash-exp",
                contents=prompt
            )
            
            if not response or not response.text:
                raise ValueError("Empty response from API")
            
            return response.text
            
        except Exception as e:
            error_str = str(e)
            
            # Check if it's a quota/rate limit error
            if "429" in error_str or "RESOURCE_EXHAUSTED" in error_str or "quota" in error_str.lower():
                # Mark this key as failed
                mark_api_key_failed(api_key)
                
                # Try to get next API key
                if allow_key_rotation:
                    next_key = get_next_api_key()
                    keys_list = st.session_state.get("api_keys_list", [])
                    
                    if next_key and next_key != api_key and len(keys_list) > 1:
                        key_number = keys_list.index(next_key) + 1 if next_key in keys_list else "?"
                        st.info(f"🔄 Switching to API key #{key_number} due to rate limit...")
                        # Retry immediately with new key
                        try:
                            client = get_client(next_key)
                            response = client.models.generate_content(
                                model="gemini-2.0-flash-exp",
                                contents=prompt
                            )
                            if response and response.text:
                                return response.text
                        except:
                            pass
                
                # Extract retry delay from error message
                import re
                delay_match = re.search(r'retry.*?(\d+\.?\d*)\s*s', error_str, re.IGNORECASE)
                
                if delay_match:
                    retry_delay = float(delay_match.group(1))
                else:
                    retry_delay = 30 * (2 ** attempt)
                
                if attempt < max_retries - 1:
                    st.warning(f"⏳ Rate limit hit. Waiting {retry_delay:.0f} seconds before retry {attempt + 2}/{max_retries}...")
                    time.sleep(retry_delay)
                    continue
                else:
                    st.error(f"❌ Rate limit exceeded after {max_retries} attempts. All API keys exhausted.")
                    raise
            else:
                # Other errors - retry with shorter delay
                if attempt < max_retries - 1:
                    wait_time = 2 * (2 ** attempt)
                    st.warning(f"⚠️ API error (attempt {attempt + 1}/{max_retries}): {error_str[:100]}... Retrying in {wait_time}s")
                    time.sleep(wait_time)
                else:
                    st.error(f"❌ API Error after {max_retries} attempts: {error_str}")
                    raise
    
    raise ValueError("Failed to get response after all retries")

def _ensure_json_payload(raw_text: str) -> Optional[Union[Dict, List]]:
    """Attempt to extract and decode JSON from the model output."""
    raw_text = raw_text.strip()
    if not raw_text:
        return None
    json_snippet = raw_text
    if not raw_text.lstrip().startswith(("{", "[")):
        try:
            start = raw_text.index("{")
            end = raw_text.rindex("}") + 1
            json_snippet = raw_text[start:end]
        except ValueError:
            return None
    try:
        return json.loads(json_snippet)
    except json.JSONDecodeError:
        return None

def _humanize_key(key: str) -> str:
    key = key.replace("_", " ").replace("-", " ")
    return key.strip().title()

def _normalize_text(value: Union[str, int, float]) -> str:
    text = str(value)
    text = text.replace("**", "")
    text = text.replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()

def _list_to_lines(items: List[Any], level: int = 0) -> List[str]:
    lines: List[str] = []
    for item in items:
        if item is None:
            continue
        if isinstance(item, list):
            lines.extend(_list_to_lines(item, level + 1))
            continue
        if isinstance(item, dict):
            dict_parts: List[str] = []
            for key, value in item.items():
                inner_text = clean_text_value(value)
                if inner_text:
                    dict_parts.append(f"{_humanize_key(str(key))}: {inner_text}")
            if dict_parts:
                prefix = "• " if level == 0 else "  " * level + "- "
                lines.append(f"{prefix}{'; '.join(dict_parts)}")
            continue
        text = _normalize_text(item)
        if not text:
            continue
        sublines = [part.strip() for part in text.split("\n") if part.strip()]
        if not sublines:
            continue
        first_prefix = "• " if level == 0 else "  " * level + "- "
        continuation_prefix = "  - " if level == 0 else "  " * (level + 1) + "- "
        lines.append(f"{first_prefix}{sublines[0]}")
        for extra in sublines[1:]:
            lines.append(f"{continuation_prefix}{extra}")
    return lines

def clean_text_value(value: Union[str, List, Dict[str, str], None]) -> str:
    if value is None:
        return ""
    if isinstance(value, list):
        bullet_lines = _list_to_lines(value)
        return "\n".join(bullet_lines)
    if isinstance(value, dict):
        parts = [f"{_humanize_key(str(k))}: {clean_text_value(v)}" for k, v in value.items()]
        return "\n".join(part for part in parts if part.strip())
    return _normalize_text(value)

BASE_FIELD_LABELS: Dict[str, str] = {
    "bid_number": "Bid Number",
    "start_date": "Start Date",
    "end_date": "End Date",
    "city": "City",
    "organisation_name": "Organization",
    "organization_name": "Organization",
    "department_name": "Department Name",
    "period_of_contract": "Period Of Contract",
    "estimated_bid_amount": "Est. Bid Amount (Rs.)",
    "category_or_type": "Service Category",
    "scope_of_work": "Scope Of Work",
    "scope": "Scope Of Work",
    "eligibility_criteria": "Eligibility Criteria",
    "delivery_or_execution_mode": "Offsite / Onsite",
    "minimum_turnover_requirement": "Minimum Turnover Requirement",
    "experience_requirement": "Experience Requirement",
    "earnest_money_deposit": "Earnest Money Deposit",
    "bid_security": "Bid Security",
    "important_dates": "Important Dates",
    "contact_details": "Contact Details",
}

def normalize_delivery_mode(raw_text: str) -> Tuple[str, Optional[str]]:
    """Collapse delivery/execution notes to Onsite/Offsite/Hybrid labels."""
    if not raw_text:
        return "Not Found", None
    cleaned = raw_text.strip()
    if cleaned.lower() in {"not found", "na", "n/a"}:
        return "Not Found", None
    lowered = cleaned.lower()
    onsite = bool(re.search(r"\bon[-\s]?site\b", lowered))
    offsite = bool(
        re.search(r"\boff[-\s]?site\b", lowered)
        or "remote" in lowered
        or "work from home" in lowered
    )
    if onsite and offsite:
        return "Hybrid", cleaned
    if onsite:
        return "Onsite", cleaned if len(cleaned) > len("Onsite") else None
    if offsite:
        return "Offsite", cleaned if len(cleaned) > len("Offsite") else None
    return "Not Found", cleaned

def parse_ai_response(response_text: str) -> Dict[str, str]:
    """Parse AI response and return only structured data without raw JSON artifacts."""
    data: Dict[str, str] = {}
    
    # Try to extract JSON
    payload = _ensure_json_payload(response_text)
    
    if not payload:
        st.warning("⚠️ Could not extract valid JSON from AI response")
        st.text_area("Raw AI Response (first 500 chars)", response_text[:500], height=150)
        # Return empty data with all base fields
        for label in BASE_FIELD_LABELS.values():
            data[label] = "Not Found"
        return data
    
    if isinstance(payload, dict):
        additional_entries: List[str] = []
        additional_fields = payload.pop("additional_fields", {})
        
        # Process only recognized base fields
        for key, value in payload.items():
            if key in BASE_FIELD_LABELS:
                label = BASE_FIELD_LABELS[key]
                text_value = clean_text_value(value)
                
                if key == "delivery_or_execution_mode":
                    normalized, detail = normalize_delivery_mode(text_value)
                    data[label] = normalized
                    if detail and detail != normalized:
                        additional_entries.append(f"Delivery Details: {detail}")
                else:
                    data[label] = text_value
        
        # Process additional_fields object
        if isinstance(additional_fields, dict):
            for key, value in additional_fields.items():
                text_value = clean_text_value(value)
                if text_value and text_value.lower() not in {"not found", "na", "n/a"}:
                    additional_entries.append(f"{_humanize_key(key)}: {text_value}")
        
        # Add additional details if any
        if additional_entries:
            data["Additional Details"] = "\n".join(additional_entries)
        
        # Fill in missing base fields
        for label in BASE_FIELD_LABELS.values():
            data.setdefault(label, "Not Found")
    else:
        st.warning(f"⚠️ AI returned unexpected data type: {type(payload)}")
        # Fallback: couldn't parse as JSON
        for label in BASE_FIELD_LABELS.values():
            data[label] = "Not Found"
    
    return data

@retry_on_failure(max_attempts=3, delay=2)
def analyze_document(
    name: str,
    data: bytes,
    prompt_template: str,
    api_key: str,
    source_label: Optional[str] = None,
    text_override: Optional[str] = None,
) -> Optional[Dict[str, str]]:
    """Analyze a single document with retry logic - FIXED VERSION."""
    suffix = Path(name).suffix.lower()
    extracted_text: Optional[str] = None

    # Extract text based on file type
    if text_override:
        extracted_text = text_override
        st.info(f"📄 Using pre-extracted text ({len(extracted_text):,} characters)")
    elif suffix == ".pdf":
        st.info(f"🔍 Extracting text from PDF: {name}")
        extracted_text = extract_text_from_pdf(data)
    elif suffix in {".doc", ".docx"}:
        st.info(f"🔍 Extracting text from DOCX: {name}")
        extracted_text = extract_text_from_docx(data)
    else:
        try:
            extracted_text = data.decode("utf-8")
            st.info(f"🔍 Decoded text file: {name}")
        except UnicodeDecodeError:
            st.warning(f"⚠️ Cannot extract text from {name}. Unsupported format.")
            return None

    # Validate text extraction
    if not extracted_text or not extracted_text.strip():
        st.error(f"❌ No text content extracted from {name}. File may be scanned/image-based PDF.")
        return None
    
    # Log text extraction success
    text_length = len(extracted_text)
    st.success(f"✅ Extracted {text_length:,} characters from {name}")
    
    # Show preview of extracted text
    preview = extracted_text[:200].replace('\n', ' ').strip()
    st.caption(f"Preview: {preview}...")

    # Run AI analysis with retry
    try:
        st.info(f"🤖 Sending to AI for analysis...")
        response_text = run_model(prompt_template, extracted_text, api_key)
        
        if not response_text or len(response_text.strip()) < 50:
            st.error(f"❌ AI returned empty/invalid response for {name}")
            st.error(f"Response length: {len(response_text) if response_text else 0} characters")
            return None
        
        # Log successful API response
        st.success(f"✅ AI analysis complete for {name} (response: {len(response_text)} chars)")
        
        # Show preview of AI response
        response_preview = response_text[:300].replace('\n', ' ').strip()
        st.caption(f"AI Response Preview: {response_preview}...")
        
    except Exception as exc:
        st.error(f"❌ AI analysis failed for {name}: {exc}")
        import traceback
        st.error(f"Full error: {traceback.format_exc()}")
        return None

    # Parse response - FIXED VERSION
    try:
        st.info(f"📊 Parsing AI response for {name}...")
        parsed_data = parse_ai_response(response_text)
        
        if not parsed_data:
            st.error(f"❌ Failed to parse AI response - got empty data")
            st.text_area("Raw AI Response", response_text[:1000], height=200)
            return None
        
        # Check if we got any meaningful data (excluding metadata fields)
        non_empty_fields = sum(
            1 for k, v in parsed_data.items() 
            if k not in ["File Name", "Source Document"] 
            and v not in ["Not Found", "", None]
        )
        
        st.info(f"📈 Found {non_empty_fields} non-empty fields in {name}")
        
        if non_empty_fields == 0:
            st.warning(f"⚠️ All fields empty for {name}")
            st.text_area(f"Raw Response for {name}", response_text[:1000], height=200)
        else:
            # Show some extracted fields
            sample_fields = [f"{k}: {str(v)[:50]}..." for k, v in list(parsed_data.items())[:3] if v not in ["Not Found", ""]]
            if sample_fields:
                st.success(f"Sample fields: {'; '.join(sample_fields)}")
        
        # Always add metadata fields
        parsed_data['File Name'] = name
        if source_label:
            parsed_data['Source Document'] = source_label
        
        return parsed_data
        
    except Exception as exc:
        st.error(f"❌ Failed to parse response for {name}: {exc}")
        import traceback
        st.error(f"Traceback: {traceback.format_exc()}")
        st.text_area("Raw Response That Failed to Parse", response_text[:1000], height=200)
        return None

def load_existing_excel(uploaded_file) -> Tuple[Optional[pd.DataFrame], Optional[str]]:
    """Load existing Excel file and return DataFrame and sheet name."""
    try:
        excel_file = pd.ExcelFile(uploaded_file)
        
        # Try to find the main data sheet
        preferred_sheets = ['Primary RFPs', 'All Documents', 'Supporting Docs']
        sheet_name = None
        
        for preferred in preferred_sheets:
            if preferred in excel_file.sheet_names:
                sheet_name = preferred
                break
        
        if not sheet_name:
            sheet_name = excel_file.sheet_names[0]
        
        df = pd.read_excel(uploaded_file, sheet_name=sheet_name)
        return df, sheet_name
    except Exception as e:
        st.error(f"Error loading Excel file: {e}")
        return None, None

def create_excel_output(
    primary_results: List[Dict], 
    supporting_results: List[Dict] = None,
    existing_df: Optional[pd.DataFrame] = None
) -> bytes:
    """Create structured Excel output with multiple sheets - with append support."""
    output = BytesIO()
    
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        # Prepare primary results
        if primary_results:
            primary_df = pd.DataFrame(primary_results).fillna("")
            primary_df = order_dataframe_columns(primary_df)
            
            # Append to existing data if provided
            if existing_df is not None:
                # Align columns
                for col in primary_df.columns:
                    if col not in existing_df.columns:
                        existing_df[col] = "Not Found"
                for col in existing_df.columns:
                    if col not in primary_df.columns:
                        primary_df[col] = "Not Found"
                
                # Reorder to match
                primary_df = primary_df[existing_df.columns]
                
                # Append
                primary_df = pd.concat([existing_df, primary_df], ignore_index=True)
            
            primary_df.to_excel(writer, index=False, sheet_name='Primary RFPs')
            sheet = writer.sheets['Primary RFPs']
            format_worksheet(sheet)
        
        # Supporting documents sheet
        if supporting_results:
            supporting_df = pd.DataFrame(supporting_results).fillna("")
            supporting_df = order_dataframe_columns(supporting_df)
            supporting_df.to_excel(writer, index=False, sheet_name='Supporting Docs')
            
            sheet = writer.sheets['Supporting Docs']
            format_worksheet(sheet)
        
        # Combined summary if both exist
        if primary_results and supporting_results:
            combined_list = primary_results.copy()
            if existing_df is not None:
                # Convert existing df back to dict list
                combined_list = existing_df.to_dict('records') + combined_list
            combined_list.extend(supporting_results)
            
            combined_df = pd.DataFrame(combined_list).fillna("")
            combined_df = order_dataframe_columns(combined_df)
            combined_df.to_excel(writer, index=False, sheet_name='All Documents')
            
            sheet = writer.sheets['All Documents']
            format_worksheet(sheet)
        
        # Metadata sheet
        total_primary = len(primary_results) if primary_results else 0
        total_supporting = len(supporting_results) if supporting_results else 0
        existing_count = len(existing_df) if existing_df is not None else 0
        
        metadata_df = pd.DataFrame({
            'Metric': [
                'Previous Records',
                'New Primary RFPs',
                'New Supporting Documents',
                'Total Documents in File',
                'Analysis Date'
            ],
            'Value': [
                existing_count,
                total_primary,
                total_supporting,
                existing_count + total_primary + total_supporting,
                pd.Timestamp.now().strftime('%Y-%m-%d %H:%M:%S')
            ]
        })
        metadata_df.to_excel(writer, index=False, sheet_name='Metadata')
    
    return output.getvalue()

def order_dataframe_columns(df: pd.DataFrame) -> pd.DataFrame:
    """Order DataFrame columns in a structured way."""
    column_structure = {
        "Document Info": ["File Name", "Source Document", "Bid Number"],
        "Timeline": ["Start Date", "End Date", "Period Of Contract", "Important Dates"],
        "Organization": ["Organization", "Department Name", "City", "Contact Details"],
        "Financial": [
            "Est. Bid Amount (Rs.)",
            "Minimum Turnover Requirement",
            "Earnest Money Deposit",
            "Bid Security"
        ],
        "Service Details": [
            "Service Category",
            "Scope Of Work",
            "Offsite / Onsite",
            "Experience Requirement"
        ],
        "Eligibility": ["Eligibility Criteria"],
        "Additional": ["Additional Details"]
    }
    
    ordered_columns = []
    for group_cols in column_structure.values():
        for col in group_cols:
            if col in df.columns and col not in ordered_columns:
                ordered_columns.append(col)
    
    for col in df.columns:
        if col not in ordered_columns:
            ordered_columns.append(col)
    
    for col in ordered_columns:
        if col not in df.columns:
            df[col] = "Not Found"
    
    return df[ordered_columns]

def format_worksheet(sheet):
    """Format worksheet with proper column widths."""
    for column in sheet.columns:
        max_length = 0
        column_letter = column[0].column_letter
        for cell in column:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = min(max_length + 2, 50)
        sheet.column_dimensions[column_letter].width = adjusted_width

# --- Sidebar Configuration ---
st.sidebar.header("Configuration")

# Load API keys
if "api_keys_list" not in st.session_state or not st.session_state["api_keys_list"]:
    st.session_state["api_keys_list"] = load_api_keys()

stored_keys = st.session_state.get("api_keys_list", [])
stored_key = stored_keys[0] if stored_keys else None

if stored_keys:
    active_count = len([k for k in stored_keys if k not in st.session_state.get("failed_api_keys", set())])
    st.sidebar.success(f"✅ {len(stored_keys)} API key(s) loaded ({active_count} active)")
    
    # Show key rotation status
    if len(stored_keys) > 1:
        current_idx = st.session_state.get("current_api_key_index", 0)
        st.sidebar.info(f"🔄 Currently using key #{current_idx + 1}")
else:
    st.sidebar.info("Enter your Gemini API key(s). Multiple keys will auto-rotate when limits are hit.")

with st.sidebar.expander("Manage API Keys", expanded=not bool(stored_keys)):
    st.write("**Add API Keys (one per line)**")
    
    # Text area for multiple keys
    default_keys_text = "\n".join(stored_keys) if stored_keys else ""
    api_keys_input = st.text_area(
        "Google Gemini API Keys",
        value=default_keys_text,
        height=150,
        help="Enter one API key per line. Keys will rotate automatically when rate limits are hit."
    )
    
    col1, col2 = st.columns(2)
    with col1:
        if st.button("💾 Save Keys"):
            keys_list = [k.strip() for k in api_keys_input.split('\n') if k.strip()]
            if keys_list:
                save_api_keys(keys_list)
    
    with col2:
        if st.button("🔄 Reset Failed Keys"):
            st.session_state["failed_api_keys"] = set()
            st.session_state["current_api_key_index"] = 0
            st.success("✅ All keys reset and available again!")
    
    # Show current status
    if stored_keys:
        st.write("---")
        st.write("**Key Status:**")
        failed_keys = st.session_state.get("failed_api_keys", set())
        for idx, key in enumerate(stored_keys, 1):
            masked_key = f"{key[:8]}...{key[-4:]}" if len(key) > 12 else "****"
            status = "❌ Rate Limited" if key in failed_keys else "✅ Active"
            current_marker = "👉 " if idx - 1 == st.session_state.get("current_api_key_index", 0) else "   "
            st.text(f"{current_marker}Key #{idx}: {masked_key} - {status}")

prompt_template = st.sidebar.text_area(
    "Prompt Template",
    value=st.session_state.get("prompt_template", DEFAULT_PROMPT),
    height=240,
)

if st.sidebar.button("Reset Prompt to Default"):
    prompt_template = DEFAULT_PROMPT

st.session_state["prompt_template"] = prompt_template

# --- EXCEL APPEND SECTION ---
st.header("📊 Excel File Management")
st.write("Upload an existing Excel file to append new analysis results, or start fresh.")

col1, col2 = st.columns([2, 1])

with col1:
    existing_excel = st.file_uploader(
        "Upload Existing Excel File (Optional)",
        type=["xlsx"],
        key="existing_excel",
        help="Upload your previous analysis Excel file. New results will be appended to it."
    )

with col2:
    if existing_excel:
        if st.button("Load Excel File"):
            df, sheet_name = load_existing_excel(existing_excel)
            if df is not None:
                st.session_state["existing_excel_data"] = df
                st.session_state["existing_excel_filename"] = existing_excel.name
                st.success(f"✅ Loaded {len(df)} existing records from '{sheet_name}' sheet")
            else:
                st.error("Failed to load Excel file")
    
    if st.session_state.get("existing_excel_data") is not None:
        if st.button("Clear Loaded Excel"):
            st.session_state["existing_excel_data"] = None
            st.session_state["existing_excel_filename"] = None
            st.rerun()

if st.session_state.get("existing_excel_data") is not None:
    st.info(f"📁 Using existing file: **{st.session_state['existing_excel_filename']}** ({len(st.session_state['existing_excel_data'])} records)")

st.write("---")

# --- STEP 1: Analyze Primary RFP Documents ---
st.header("Step 1: Analyze Primary RFP Documents")
st.write("Upload your main RFP PDF files and analyze them to get structured data.")

uploaded_files = st.file_uploader(
    "Choose RFP PDF files",
    type=["pdf"],
    accept_multiple_files=True,
    key="primary_upload"
)

if st.button("Analyze Primary RFPs", type="primary", disabled=not uploaded_files):
    # Get active API key
    active_api_key = get_next_api_key()
    
    if not active_api_key:
        st.warning("⚠️ Please provide at least one Gemini API key in the sidebar.")
    else:
        st.session_state["primary_analysis_results"] = []
        st.session_state["extracted_links"] = {}
        st.session_state["attachment_cache"] = {}
        
        progress_bar = st.progress(0)
        status_text = st.empty()
        
        # Create a detailed log container
        with st.expander("📋 Detailed Analysis Log", expanded=True):
            log_container = st.container()
        
        total_files = len(uploaded_files)
        successful = 0
        failed = 0
        
        for idx, pdf_file in enumerate(uploaded_files):
            with log_container:
                st.write(f"\n**{'='*60}**")
                st.write(f"**Processing {idx+1}/{total_files}: {pdf_file.name}**")
                st.write(f"**{'='*60}**")
            
            status_text.text(f"Processing {idx+1}/{total_files}: {pdf_file.name}...")
            
            try:
                # Read PDF bytes - FRESH for each file
                pdf_bytes = pdf_file.getvalue()
                
                # Extract text FRESH for each document
                with log_container:
                    st.info(f"🔍 Extracting text from {pdf_file.name}...")
                
                extracted_text = extract_text_from_pdf(pdf_bytes)
                
                if not extracted_text:
                    with log_container:
                        st.error(f"⚠️ Could not extract text from {pdf_file.name}. Skipping.")
                    failed += 1
                    progress_bar.progress((idx + 1) / total_files)
                    continue
                
                with log_container:
                    st.success(f"✅ Extracted {len(extracted_text):,} characters from {pdf_file.name}")
                
                # Add a small delay between API calls to avoid rate limiting
                if idx > 0:
                    time.sleep(2)  # 2 second delay between documents to help avoid rate limits
                
                # Analyze the document
                status_text.text(f"Analyzing {idx+1}/{total_files}: {pdf_file.name}...")
                
                # Get current active key (may have rotated)
                active_api_key = get_next_api_key()
                
                result = analyze_document(
                    pdf_file.name,
                    pdf_bytes,
                    prompt_template,
                    active_api_key,
                    source_label="Primary RFP",
                    text_override=extracted_text,
                )
                
                if result:
                    # Verify the result has meaningful data
                    non_empty_fields = sum(1 for k, v in result.items() 
                                          if k not in ["File Name", "Source Document"] 
                                          and v not in ["Not Found", "", None])
                    
                    if non_empty_fields > 0:
                        st.session_state["primary_analysis_results"].append(result)
                        with log_container:
                            st.success(f"✅ Successfully analyzed {pdf_file.name} ({non_empty_fields} fields extracted)")
                        successful += 1
                    else:
                        with log_container:
                            st.warning(f"⚠️ Analysis returned no data for {pdf_file.name}")
                        failed += 1
                else:
                    with log_container:
                        st.error(f"⚠️ Failed to analyze {pdf_file.name}")
                    failed += 1
                
                # Extract links and attachments for Step 2
                embedded_links = extract_embedded_links_from_pdf(pdf_bytes)
                text_links = find_reference_links(extracted_text)
                
                all_links = []
                for url, link_text, page_info, label in embedded_links:
                    category = categorize_link(url, link_text, label)
                    all_links.append((url, category, f"{link_text} ({page_info})", label))
                
                for url, category in text_links:
                    if not any(link[0] == url for link in all_links):
                        all_links.append((url, category, "Plain text URL", ""))
                
                if all_links:
                    st.session_state["extracted_links"][pdf_file.name] = all_links
                    with log_container:
                        st.info(f"🔗 Found {len(all_links)} links in {pdf_file.name}")
                
                # Extract attachments
                attachments = extract_attachments_from_pdf(pdf_bytes)
                if attachments:
                    st.session_state["attachment_cache"][pdf_file.name] = attachments
                    with log_container:
                        st.info(f"📎 Found {len(attachments)} attachments in {pdf_file.name}")
                
            except Exception as e:
                with log_container:
                    st.error(f"❌ Unexpected error processing {pdf_file.name}: {str(e)}")
                failed += 1
            
            progress_bar.progress((idx + 1) / total_files)
        
        # Final summary
        status_text.text("✅ Analysis complete!")
        
        with log_container:
            st.write(f"\n**{'='*60}**")
            st.write("**FINAL SUMMARY**")
            st.write(f"**{'='*60}**")
            st.write(f"✅ **Successful:** {successful}/{total_files}")
            st.write(f"❌ **Failed:** {failed}/{total_files}")
            st.write(f"📊 **Success Rate:** {(successful/total_files*100):.1f}%")
        
        st.session_state["current_step"] = 1
        
        if successful > 0:
            st.balloons()
            st.success(f"Successfully analyzed {successful} out of {total_files} documents!")
        else:
            st.error("No documents were successfully analyzed. Please check the logs above.")
        
        st.rerun()

# Display Step 1 Results
if st.session_state["primary_analysis_results"]:
    st.write("---")
    st.subheader("Primary RFP Analysis Results")
    
    primary_df = pd.DataFrame(st.session_state["primary_analysis_results"]).fillna("")
    primary_df = order_dataframe_columns(primary_df)
    
    st.dataframe(primary_df, use_container_width=True)
    
    # Download primary results
    existing_data = st.session_state.get("existing_excel_data")
    excel_data = create_excel_output(
        st.session_state["primary_analysis_results"],
        existing_df=existing_data
    )
    
    col1, col2 = st.columns(2)
    with col1:
        filename_suffix = "appended" if existing_data is not None else "new"
        st.download_button(
            label=f"📥 Download Primary RFP Analysis (Excel) - {filename_suffix.upper()}",
            data=excel_data,
            file_name=f"primary_rfp_analysis_{filename_suffix}_{pd.Timestamp.now().strftime('%Y%m%d_%H%M%S')}.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
    
    with col2:
        csv_data = primary_df.to_csv(index=False)
        st.download_button(
            label="📥 Download Primary RFP Analysis (CSV)",
            data=csv_data,
            file_name=f"primary_rfp_analysis_{pd.Timestamp.now().strftime('%Y%m%d_%H%M%S')}.csv",
            mime="text/csv",
        )
    
    # Show summary of available supporting documents
    total_links = sum(len(links) for links in st.session_state.get("extracted_links", {}).values())
    total_attachments = sum(len(atts) for atts in st.session_state.get("attachment_cache", {}).values())
    
    if total_links > 0 or total_attachments > 0:
        st.info(f"📎 Found **{total_links}** document links and **{total_attachments}** embedded attachments. Proceed to Step 2 to analyze them.")

# --- STEP 2: Analyze Supporting Documents ---
if st.session_state["primary_analysis_results"]:
    st.write("---")
    st.header("Step 2: Analyze Supporting Documents (Optional)")
    st.write("Download and analyze ATC, SOW, and other supporting documents referenced in your RFPs.")
    
    # Section 2A: Embedded Attachments
    attachment_cache = st.session_state.get("attachment_cache", {})
    if attachment_cache:
        st.subheader("2A: Embedded Attachments")
        st.write("These files are embedded directly in the PDF:")
        
        for pdf_name, attachments in attachment_cache.items():
            with st.expander(f"Attachments in {pdf_name}", expanded=True):
                for idx, (att_name, att_data, desc) in enumerate(attachments, start=1):
                    col1, col2, col3 = st.columns([2, 2, 2])
                    
                    with col1:
                        st.write(f"**{att_name}**")
                        if desc:
                            st.caption(desc)
                    
                    with col2:
                        file_size = len(att_data) / 1024
                        st.caption(f"Size: {file_size:.1f} KB")
                    
                    with col3:
                        mime_type = mimetypes.guess_type(att_name)[0] or "application/octet-stream"
                        
                        col3a, col3b = st.columns(2)
                        with col3a:
                            st.download_button(
                                label="Download",
                                data=att_data,
                                file_name=att_name,
                                mime=mime_type,
                                key=f"dl_att_{pdf_name}_{idx}",
                            )
                        
                        with col3b:
                            if st.button("Analyze", key=f"analyze_att_{pdf_name}_{idx}"):
                                active_api_key = get_next_api_key()
                                if active_api_key:
                                    with st.spinner(f"Analyzing {att_name}..."):
                                        result = analyze_document(
                                            att_name,
                                            att_data,
                                            prompt_template,
                                            active_api_key,
                                            source_label=f"Attachment from {pdf_name}",
                                        )
                                        if result:
                                            st.session_state["supporting_analysis_results"].append(result)
                                            st.success(f"✅ Analyzed {att_name}")
                                            st.rerun()
                                        else:
                                            st.error(f"❌ Failed to analyze {att_name}")
    
    # Section 2B: External Links
    extracted_links = st.session_state.get("extracted_links", {})
    if extracted_links:
        st.write("---")
        st.subheader("2B: External Document Links")
        st.write("Select documents to download and analyze:")
        
        # Temporary storage for selected links and downloaded files
        if "temp_selected_links" not in st.session_state:
            st.session_state["temp_selected_links"] = {}
        if "downloaded_files" not in st.session_state:
            st.session_state["downloaded_files"] = {}
        
        for pdf_name, links in extracted_links.items():
            with st.expander(f"Links from {pdf_name}", expanded=True):
                if pdf_name not in st.session_state["temp_selected_links"]:
                    st.session_state["temp_selected_links"][pdf_name] = []
                
                for idx, link_data in enumerate(links):
                    if len(link_data) >= 4:
                        url, category, context, label = link_data[:4]
                    elif len(link_data) == 3:
                        url, category, context = link_data
                        label = ""
                    else:
                        url, category = link_data
                        context = ""
                        label = ""
                    
                    col1, col2, col3, col4 = st.columns([0.5, 3, 1.5, 1])
                    
                    with col1:
                        is_selected = st.checkbox(
                            "Select",
                            key=f"select_link_{pdf_name}_{idx}",
                            value=url in st.session_state["temp_selected_links"][pdf_name]
                        )
                        
                        if is_selected and url not in st.session_state["temp_selected_links"][pdf_name]:
                            st.session_state["temp_selected_links"][pdf_name].append(url)
                        elif not is_selected and url in st.session_state["temp_selected_links"][pdf_name]:
                            st.session_state["temp_selected_links"][pdf_name].remove(url)
                    
                    with col2:
                        if label and label != "Document":
                            st.markdown(f"**{label}**")
                        st.markdown(f"[{url}]({url})")
                        if context:
                            st.caption(context)
                    
                    with col3:
                        if category == "ATC Document":
                            st.markdown("🔷 ATC")
                        elif category == "SOW Document":
                            st.markdown("📋 SOW")
                        elif category == "Corrigendum":
                            st.markdown("📝 Corrigendum")
                        elif category == "Annexure":
                            st.markdown("📎 Annexure")
                        elif category == "Related Document":
                            st.markdown("📄 Related")
                        else:
                            st.markdown(category)
                    
                    with col4:
                        # Show download button if file has been fetched
                        file_key = f"{pdf_name}||{url}"
                        if file_key in st.session_state["downloaded_files"]:
                            file_info = st.session_state["downloaded_files"][file_key]
                            st.download_button(
                                label="Download",
                                data=file_info["bytes"],
                                file_name=file_info["name"],
                                mime=file_info.get("mime", "application/octet-stream"),
                                key=f"dl_ext_{pdf_name}_{idx}",
                            )
        
        # Analyze selected button
        st.write("---")
        total_selected = sum(len(links) for links in st.session_state["temp_selected_links"].values())
        
        if st.button(f"Download & Analyze Selected Documents ({total_selected})", type="primary", disabled=total_selected == 0):
            active_api_key = get_next_api_key()
            
            if total_selected == 0:
                st.warning("Please select at least one document to download.")
            elif not active_api_key:
                st.warning("⚠️ Please provide at least one API key in the sidebar.")
            else:
                progress_bar = st.progress(0)
                status_text = st.empty()
                
                processed = 0
                for pdf_name, selected_urls in st.session_state["temp_selected_links"].items():
                    for url in selected_urls:
                        status_text.text(f"Processing {processed+1}/{total_selected}: {url.split('/')[-1]}...")
                        
                        fetched = fetch_external_document(url)
                        if fetched:
                            doc_name, doc_bytes, content_type = fetched
                            
                            # Store the downloaded file for later download
                            file_key = f"{pdf_name}||{url}"
                            st.session_state["downloaded_files"][file_key] = {
                                "name": doc_name,
                                "bytes": doc_bytes,
                                "mime": content_type or mimetypes.guess_type(doc_name)[0]
                            }
                            
                            # Analyze the document
                            active_api_key = get_next_api_key()  # Get current active key
                            result = analyze_document(
                                doc_name,
                                doc_bytes,
                                prompt_template,
                                active_api_key,
                                source_label=f"Downloaded from {pdf_name}",
                            )
                            
                            if result:
                                st.session_state["supporting_analysis_results"].append(result)
                                st.success(f"✅ Analyzed {doc_name}")
                            else:
                                st.warning(f"⚠️ Failed to analyze {doc_name}")
                        else:
                            st.error(f"❌ Failed to download from {url}")
                        
                        processed += 1
                        progress_bar.progress(processed / total_selected)
                
                status_text.text("✅ Analysis complete!")
                st.balloons()
                st.rerun()
    
    # Display Step 2 Results
    if st.session_state.get("supporting_analysis_results"):
        st.write("---")
        st.subheader("Supporting Documents Analysis Results")
        
        supporting_df = pd.DataFrame(st.session_state["supporting_analysis_results"]).fillna("")
        supporting_df = order_dataframe_columns(supporting_df)
        
        st.dataframe(supporting_df, use_container_width=True)
        
        # Download combined results
        st.write("---")
        st.subheader("📥 Download Complete Analysis")
        
        existing_data = st.session_state.get("existing_excel_data")
        combined_excel = create_excel_output(
            st.session_state["primary_analysis_results"],
            st.session_state["supporting_analysis_results"],
            existing_df=existing_data
        )
        
        col1, col2, col3 = st.columns(3)
        
        with col1:
            filename_suffix = "appended" if existing_data is not None else "new"
            st.download_button(
                label=f"📥 Combined Analysis (Excel) - {filename_suffix.upper()}",
                data=combined_excel,
                file_name=f"complete_rfp_analysis_{filename_suffix}_{pd.Timestamp.now().strftime('%Y%m%d_%H%M%S')}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        
        with col2:
            # Supporting docs only
            supporting_excel = create_excel_output([], st.session_state["supporting_analysis_results"])
            st.download_button(
                label="📥 Supporting Docs Only (Excel)",
                data=supporting_excel,
                file_name=f"supporting_docs_{pd.Timestamp.now().strftime('%Y%m%d_%H%M%S')}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        
        with col3:
            # Combined CSV
            all_results = st.session_state["primary_analysis_results"] + st.session_state["supporting_analysis_results"]
            combined_df = pd.DataFrame(all_results).fillna("")
            combined_df = order_dataframe_columns(combined_df)
            csv_data = combined_df.to_csv(index=False)
            st.download_button(
                label="📥 Combined Analysis (CSV)",
                data=csv_data,
                file_name=f"complete_rfp_analysis_{pd.Timestamp.now().strftime('%Y%m%d_%H%M%S')}.csv",
                mime="text/csv",
            )

# Reset button
if st.session_state["primary_analysis_results"] or st.session_state.get("supporting_analysis_results"):
    st.write("---")
    if st.button("🔄 Start New Analysis (Keep Existing Excel)"):
        st.session_state["primary_analysis_results"] = []
        st.session_state["supporting_analysis_results"] = []
        st.session_state["extracted_links"] = {}
        st.session_state["attachment_cache"] = {}
        st.session_state["temp_selected_links"] = {}
        st.session_state["downloaded_files"] = {}
        st.session_state["current_step"] = 1
        st.rerun()
    
    if st.button("🗑️ Clear Everything (Including Excel)"):
        for key in list(st.session_state.keys()):
            del st.session_state[key]
        st.rerun()

# Footer
st.write("---")
st.caption("💡 Tip: Upload an existing Excel file at the top to append new analysis results to your previous data.")
